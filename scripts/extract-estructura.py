#!/usr/bin/env python3
"""
extract-estructura.py — Convierte el Excel fuente en server/data/estructura-sgeo.json.

Usa dos hojas:
  · "Comparativa SGEO FINAL" -> elementos, subelementos, pesos, fase PDCA y referencias
  · "CCPS"                    -> catálogo de pilares/elementos del modelo CCPS RBPS

Uso:
    pip install openpyxl
    python3 scripts/extract-estructura.py [ruta_al_xlsx]

Mapeo de "Comparativa SGEO FINAL" (fila 4 = encabezados, filas 5-31 = datos):
    A  -> pdca            fase del ciclo: PLANIFICAR / HACER / VERIFICAR / ACTUAR (por subelemento)
    B  -> elemento        E1..E5 ("código: nombre"); celdas combinadas
    C  -> se              número de subelemento (1..27)
    D  -> detalle         nombre del subelemento
    E-K -> referencias    YPF 2.0, ISO 9001/14001/45001, CCPS, ISO 39001, ISO 50001
    L  -> puntajeElemento peso del elemento (15/15/15/35/20; suma 100), sólo primera fila
    M  -> ponderación     siempre 1
"""
import json
import re
import sys
import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl. Ejecute: pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "scripts" / "fuente-comparativa-sgeo.xlsx"
OUT = ROOT / "server" / "data" / "estructura-sgeo.json"
FIRST_ROW, LAST_ROW = 5, 31


def clean(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):   # celdas de referencia mal interpretadas como fecha
        return f"{v.month}.{v.day}"
    return str(v).strip().rstrip(",").strip()


def parse_comparativa(ws):
    elements, order = {}, []
    cur_pdca = cur_elem = None
    for r in range(FIRST_ROW, LAST_ROW + 1):
        cur_pdca = ws.cell(r, 1).value or cur_pdca
        cur_elem = ws.cell(r, 2).value or cur_elem
        elem = cur_elem.strip()
        code, name = (elem.split(":", 1) + [""])[:2]
        code, name = code.strip(), name.strip()
        if elem not in elements:
            elements[elem] = {"codigo": code, "nombre": name,
                              "puntajeElemento": ws.cell(r, 12).value, "subelementos": []}
            order.append(elem)
        elements[elem]["subelementos"].append({
            "se": ws.cell(r, 3).value,
            "detalle": clean(ws.cell(r, 4).value),
            "pdca": cur_pdca.strip(),
            "subElementoYPF20": clean(ws.cell(r, 5).value),
            "iso9001": clean(ws.cell(r, 6).value),
            "iso14001": clean(ws.cell(r, 7).value),
            "iso45001": clean(ws.cell(r, 8).value),
            "ccps": clean(ws.cell(r, 9).value),
            "iso39001": clean(ws.cell(r, 10).value),
            "iso50001": clean(ws.cell(r, 11).value),
            "ponderacionInicial": ws.cell(r, 13).value,
        })
    # fase del elemento = fases únicas de sus subelementos
    for e in elements.values():
        phs = []
        for s in e["subelementos"]:
            if s["pdca"] not in phs:
                phs.append(s["pdca"])
        e["pdca"] = " / ".join(phs)
    return [elements[e] for e in order]


def parse_ccps(ws):
    pilares, cur = [], None
    for r in range(1, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(a, str):
            m = re.match(r"^\s*Pilar\s+(\d+)\.\s*(.+)", a.strip(), re.I)
            if m:
                cur = {"nro": int(m.group(1)), "titulo": m.group(2).strip(),
                       "descripcion": "", "elementos": []}
                pilares.append(cur)
                continue
            if cur is not None and a.strip() and not cur["descripcion"]:
                cur["descripcion"] = a.strip()
        if isinstance(b, str) and cur is not None:
            m = re.match(r"^\s*(\d+\.\d+)\.?\s*(.+)", b.strip())
            if m:
                titulo, _, desc = m.group(2).strip().partition(":")
                cur["elementos"].append({"id": m.group(1), "titulo": titulo.strip(),
                                         "descripcion": desc.strip()})
    return pilares


def main():
    wb = openpyxl.load_workbook(SRC, data_only=False)
    elementos = parse_comparativa(wb["Comparativa SGEO FINAL"])
    pilares = parse_ccps(wb["CCPS"]) if "CCPS" in wb.sheetnames else []

    data = {
        "fuente": "Comparativa SGEO FINAL",
        "modelo": "Modelo de Excelencia Operacional YPF / Modelo YPF 2.0",
        "escalaValoracion": {"C": 1.0, "CP": 0.3, "NC": 0.0, "NA": None},
        "fasesPDCA": ["PLANIFICAR", "HACER", "VERIFICAR", "ACTUAR"],
        "umbralesEstado": [
            {"min": 0.91, "estado": "Mantener"},
            {"min": 0.81, "estado": "Optimizar"},
            {"min": 0.61, "estado": "Mejorar"},
            {"min": 0.41, "estado": "Implementar"},
            {"min": 0.0, "estado": "Crítico"},
        ],
        "elementos": elementos,
        "modeloCCPS": {
            "nombre": "CCPS RBPS — Pilares y Elementos de la Seguridad de Procesos (PSM)",
            "pilares": pilares,
        },
    }
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=2))
    total = sum(e["puntajeElemento"] for e in elementos)
    nsub = sum(len(e["subelementos"]) for e in elementos)
    print(f"OK  {OUT}")
    print(f"    {len(elementos)} elementos · {nsub} subelementos · suma de pesos L = {total}")
    print(f"    CCPS: {len(pilares)} pilares, {sum(len(p['elementos']) for p in pilares)} elementos")


if __name__ == "__main__":
    main()
